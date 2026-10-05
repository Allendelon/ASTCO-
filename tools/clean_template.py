"""Produce assets/templates/ipc-template.xlsx from the client's original IPC workbook.

Usage: python3 tools/clean_template.py "IPC - 1.xlsx"   (requires: pip install openpyxl)

What it does (layout, fonts, borders, merges and print setup are preserved):
  * removes ~19,600 legacy defined names and 9 external workbook links
    (these triggered Excel's "update links" security prompt and bloated the file to 700 KB)
  * trims the sheet to the A1:M62 print area (drops stray #DIV/0! cells outside it)
  * widens THIS PERIOD / CUMULATIVE (K, L) from 13 to 18 so 9-digit SAR amounts don't show ####,
    narrowing Remarks (M) to keep the same page width
  * fixes two label typos ("Finance Directoe", "Liquidate Damage N/A")
Then run tools/embed_template.py to refresh the base64 copy used by the browser.
"""
import sys
import openpyxl
from openpyxl.utils import column_index_from_string

src = sys.argv[1]
dst = sys.argv[2] if len(sys.argv) > 2 else "assets/templates/ipc-template.xlsx"

wb = openpyxl.load_workbook(src)
ws = wb.active

for name in list(wb.defined_names.keys()):
    del wb.defined_names[name]
wb._external_links = []

for m in list(ws.merged_cells.ranges):
    if m.min_col > 13 or m.min_row > 62:
        ws.unmerge_cells(str(m))
ws.delete_cols(14, ws.max_column)
ws.delete_rows(63, ws.max_row)
for col in list(ws.column_dimensions.keys()):
    if column_index_from_string(col) > 13:
        del ws.column_dimensions[col]

ws.column_dimensions["K"].width = 18
ws.column_dimensions["L"].width = 18
ws.column_dimensions["M"].width = 31.3
ws["D60"] = "Finance Director"
ws["C34"] = "Liquidated Damages"
ws.print_area = "A1:M62"

wb.save(dst)
print("wrote", dst)
